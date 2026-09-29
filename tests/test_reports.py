import json
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from tracker.config import PROJECT_ROOT, load_config
from tracker.db import connect
from tracker.reports import build_report, render, telegram_text, write_report

NOW = datetime(2026, 9, 28, 2, 0, tzinfo=timezone.utc)
ROWS = [
    # key, title, category, sectors, states, score, band, priority, story, outlets
    ("a", "SC directs CBSE to extend third-language exemption to Class 6", "court", ["school_education"], [], 100,
     "core", "high", "a", 6),
    ("a2", "Supreme Court asks CBSE to extend exemption to Class 6", "court", ["school_education"], [], 90,
     "core", "high", "a", 6),
    ("b", "Haryana govt transfers 3 IAS officers", "transfer", [], ["HR"], 70, "relevant", "medium", "b", 2),
    ("c", "SEC announces municipal polls schedule in Pune", "election", [], ["MH"], 73, "relevant", "medium", "c", 1),
    ("d", "UGC notifies draft rules for foreign campuses", "policy", ["higher_education"], [], 95, "core", "high",
     "d", 3),
    ("e", "Cricket final tonight", "news", [], [], 5, "not_relevant", "low", "e", 1),
]


@pytest.fixture
def cfg(tmp_path, monkeypatch):
    config = load_config(PROJECT_ROOT / "config")
    monkeypatch.setattr(type(config), "db_path", property(lambda self: tmp_path / "tracker.db"))
    conn = connect(tmp_path / "tracker.db")
    for key, title, cat, sectors, states, score, band, prio, story, outlets in ROWS:
        conn.execute(
            "INSERT INTO items(item_key, url, title, source_name, first_seen, category, sectors, states, final_score, "
            "band, priority, story_id, sources_count, reasons, updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (key, f"https://e.com/{key}", title, "Outlet", "2026-09-27T20:00:00Z", cat, json.dumps(sectors),
             json.dumps(states), score, band, prio, story, outlets, json.dumps(["rule +40"]), "2026-09-27T20:00:00Z"))
    conn.execute("INSERT INTO runs(started_at, finished_at, status, errors) VALUES(?,?,?,?)",
                 ("2026-09-27T21:00:00Z", "2026-09-27T21:02:00Z", "ok", json.dumps({"gn-transfers-ta": "timeout"})))
    conn.commit()
    conn.close()
    return config


def test_report_groups_stories_and_trackers(cfg):
    data = build_report(cfg, "daily", now=NOW)
    titles = [s.title for s in data.stories]
    assert "Cricket final tonight" not in titles
    assert titles.count("Supreme Court asks CBSE to extend exemption to Class 6") == 0  # same story, lower score
    assert data.stats["stories"] == 4 and data.stats["high_priority"] == 2
    assert data.top[0].priority == "high"
    assert [s.item_key for s in data.trackers["transfer"]] == ["b"]
    sectors = [label for _, label, _ in data.by_sector]
    assert sectors[0] == "School education" and sectors[-1].startswith("Governance")
    assert data.health == {"runs": 1, "ok": 1, "failed_sources": ["gn-transfers-ta"], "last_run": "2026-09-27T21:02:00Z"}


def test_report_files_render(cfg, tmp_path):
    data = build_report(cfg, "weekly", now=NOW)
    files = write_report(cfg, data, out_root=tmp_path / "reports")
    digest = files["digest"].read_text(encoding="utf-8")
    assert "SC directs CBSE" in digest and "reported by 6 outlets" in digest and "Haryana" in digest
    brief = render(data, "brief.html")
    assert "Top developments" in brief and "Transfers &amp; postings" in brief
    wb = load_workbook(files["excel"])
    assert wb.sheetnames == ["Summary", "All stories", "Transfers", "Elections", "Court rulings", "Policy"]
    assert wb["All stories"].max_row == 5
    assert "high priority" in telegram_text(data)


def test_top_developments_skip_other_outlets_takes_on_same_event(cfg):
    from tracker.reports import pick_top

    data = build_report(cfg, "daily", now=NOW)
    twin = data.stories[0].__class__(**{**data.stories[0].__dict__, "story_id": "z", "item_key": "z",
                                        "title": "CBSE third-language exemption: SC directs board to extend it to Class 6"})
    top = pick_top([*data.stories, twin], 8)
    assert sum("third" in s.title.lower() for s in top) == 1


def test_telegram_caption_fits_the_limit(cfg):
    from tracker.reports import telegram_caption

    data = build_report(cfg, "daily", now=NOW)
    caption = telegram_caption(data)
    assert len(caption) <= 1024 and "Daily brief" in caption and "PDF" in caption


def test_pdf_rendering_when_a_browser_exists(tmp_path):
    from tracker.pdf import find_browser, html_to_pdf

    if not find_browser():
        return  # no Chromium on this machine; the report falls back to a text message
    html = tmp_path / "x.html"
    html.write_text("<html><body><h1>मुंबई — test</h1></body></html>", encoding="utf-8")
    pdf = html_to_pdf(html, tmp_path / "x.pdf")
    assert pdf and pdf.read_bytes()[:4] == b"%PDF"
