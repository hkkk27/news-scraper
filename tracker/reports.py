"""Daily and weekly reports: analyst digest, executive brief, Excel workbook, Telegram brief, email.

Everything is built from the items table, grouped into stories (one card per story, with the
number of outlets that reported it). The same data feeds every format, so they always agree.
"""

from __future__ import annotations

import json
import re
import smtplib
import sqlite3
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path

from tracker.config import AppConfig
from tracker.log import get_logger
from tracker.taxonomy import Taxonomy, load_taxonomy

log = get_logger("tracker.reports")

BAND_ORDER = ["core", "relevant", "peripheral", "not_relevant"]
PRIORITY_RANK = {"high": 0, "medium": 1, "low": 2}
IST = timezone(timedelta(hours=5, minutes=30))
GOVERNANCE = "governance"  # pseudo-sector for transfers/elections/politics without a sector link


@dataclass
class Story:
    story_id: str
    item_key: str
    title: str
    url: str
    source_name: str
    outlets: int
    score: int
    band: str
    priority: str
    category: str
    category_label: str
    sectors: list[str]
    sector_labels: list[str]
    states: list[str]
    state_names: list[str]
    level: str
    language: str
    first_seen: str
    decided_by: str
    reasons: list[str] = field(default_factory=list)

    @property
    def where(self) -> str:
        return ", ".join(self.state_names) if self.state_names else "National"

    @property
    def is_link(self) -> bool:
        return self.url.startswith("http")


@dataclass
class ReportData:
    kind: str                     # daily | weekly
    date: str                     # report date in IST (YYYY-MM-DD)
    title: str
    period_label: str
    generated_at: str
    since: str
    stories: list[Story]
    stats: dict
    top: list[Story]
    by_sector: list[tuple[str, str, list[tuple[str, list[Story]]]]]  # (sector id, label, [(state, stories)])
    trackers: dict[str, list[Story]]
    health: dict


def _since(kind: str, now: datetime) -> datetime:
    return now - (timedelta(days=7) if kind == "weekly" else timedelta(hours=24))


def load_stories(conn: sqlite3.Connection, since: datetime, min_band: str, tax: Taxonomy) -> list[Story]:
    allowed = BAND_ORDER[: BAND_ORDER.index(min_band) + 1]
    rows = conn.execute(
        f"SELECT * FROM items WHERE first_seen >= ? AND band IN ({','.join('?' * len(allowed))}) "
        "ORDER BY final_score DESC, sources_count DESC, first_seen ASC",
        (since.strftime("%Y-%m-%dT%H:%M:%SZ"), *allowed)).fetchall()
    seen: set[str] = set()
    stories = []
    for row in rows:
        sid = row["story_id"] or row["item_key"]
        if sid in seen:  # first row per story is its best-scored item
            continue
        seen.add(sid)
        sectors = json.loads(row["sectors"] or "[]")
        states = json.loads(row["states"] or "[]")
        stories.append(Story(
            story_id=sid, item_key=row["item_key"], title=row["title"], url=row["url"],
            source_name=row["source_name"] or row["feed_id"], outlets=row["sources_count"] or 1,
            score=row["final_score"], band=row["band"], priority=row["priority"], category=row["category"],
            category_label=tax.category_label(row["category"]), sectors=sectors,
            sector_labels=[tax.sector_label(s) for s in sectors], states=states,
            state_names=[tax.state_name(s) for s in states], level=row["level"], language=row["language"],
            first_seen=row["first_seen"], decided_by=row["decided_by"],
            reasons=json.loads(row["reasons"] or "[]")))
    return stories


def _rank(story: Story) -> tuple:
    return PRIORITY_RANK.get(story.priority, 3), -story.score, -story.outlets


STOPWORDS = {"the", "a", "an", "to", "of", "in", "for", "on", "and", "from", "with", "by", "at", "as", "is", "its",
             "after", "new", "news", "live", "updates", "students", "class", "news", "report"}


def _words(title: str) -> set[str]:
    from tracker.stories import normalize_title

    return {w for w in re.findall(r"\w+", normalize_title(title))
            if w not in STOPWORDS and (len(w) > 2 or w.isdigit())}


def _same_event(a: set[str], b: set[str], overlap: float) -> bool:
    shared = len(a & b)
    return shared >= 3 and shared / max(1, min(len(a), len(b))) >= overlap


def pick_top(stories: list[Story], n: int, per_category: int = 3, overlap: float = 0.4) -> list[Story]:
    """Top developments for leadership: priority, then score, then corroboration.

    Kept varied: at most `per_category` per category, and a story is skipped when it shares ≥ 3
    key words and ≥ `overlap` of the shorter headline's words with one already chosen — another
    outlet's take on the same event.
    """
    chosen, counts, chosen_words = [], Counter(), []
    for story in sorted(stories, key=_rank):
        if story.band not in ("core", "relevant") or counts[story.category] >= per_category:
            continue
        words = _words(story.title)
        if any(_same_event(words, w, overlap) for w in chosen_words):
            continue
        chosen.append(story)
        chosen_words.append(words)
        counts[story.category] += 1
        if len(chosen) == n:
            break
    return chosen


def group_by_sector(stories: list[Story], tax: Taxonomy):
    groups: dict[str, dict[str, list[Story]]] = defaultdict(lambda: defaultdict(list))
    for story in stories:
        for sector in (story.sectors[:1] or [GOVERNANCE]):  # file each story once, under its main sector
            groups[sector][story.where].append(story)
    ordered = []
    for sector_id in [*tax.sectors, GOVERNANCE]:
        if sector_id not in groups:
            continue
        label = "Governance (transfers, elections, politics)" if sector_id == GOVERNANCE else tax.sector_label(sector_id)
        states = sorted(groups[sector_id].items(), key=lambda kv: (kv[0] != "National", -len(kv[1]), kv[0]))
        ordered.append((sector_id, label, [(state, sorted(s, key=_rank)) for state, s in states]))
    return ordered


def run_health(conn: sqlite3.Connection, since: datetime) -> dict:
    rows = conn.execute("SELECT status, errors, finished_at FROM runs WHERE started_at >= ? ORDER BY id",
                        (since.strftime("%Y-%m-%dT%H:%M:%SZ"),)).fetchall()
    ok = sum(1 for r in rows if r["status"] == "ok")
    last_errors = json.loads(rows[-1]["errors"] or "{}") if rows else {}
    return {"runs": len(rows), "ok": ok, "failed_sources": sorted(last_errors)[:12],
            "last_run": rows[-1]["finished_at"] if rows else None}


def build_report(cfg: AppConfig, kind: str = "daily", now: datetime | None = None,
                 conn: sqlite3.Connection | None = None) -> ReportData:
    from tracker.db import connect

    now = now or datetime.now(timezone.utc)
    since = _since(kind, now)
    tax = load_taxonomy(cfg.config_dir)
    own = conn is None
    conn = conn or connect(cfg.db_path)
    try:
        stories = load_stories(conn, since, cfg.settings.reports.digest_min_band, tax)
        all_rows = conn.execute("SELECT band, COUNT(*) AS n FROM items WHERE first_seen >= ? GROUP BY band",
                                (since.strftime("%Y-%m-%dT%H:%M:%SZ"),)).fetchall()
        health = run_health(conn, since)
    finally:
        if own:
            conn.close()

    bands = {r["band"]: r["n"] for r in all_rows}
    sector_counts = Counter(l for s in stories for l in (s.sector_labels or ["Governance"]))
    state_counts = Counter(n for s in stories for n in (s.state_names or ["National"]))
    category_counts = Counter(s.category_label for s in stories)
    stats = {
        "items_seen": sum(bands.values()),
        "stories": len(stories),
        "core": sum(1 for s in stories if s.band == "core"),
        "relevant": sum(1 for s in stories if s.band == "relevant"),
        "high_priority": sum(1 for s in stories if s.priority == "high"),
        "bands": bands,
        "by_sector": sector_counts.most_common(),
        "by_state": state_counts.most_common(12),
        "by_category": category_counts.most_common(),
        "languages": Counter(s.language for s in stories).most_common(),
    }
    trackers = {
        "transfer": sorted([s for s in stories if s.category == "transfer"], key=_rank)[:15],
        "election": sorted([s for s in stories if s.category == "election"], key=_rank)[:15],
        "court": sorted([s for s in stories if s.category == "court"], key=_rank)[:15],
        "policy": sorted([s for s in stories if s.category == "policy"], key=_rank)[:15],
    }
    local = now.astimezone(IST)
    period = (f"{(local - timedelta(days=7)):%d %b} – {local:%d %b %Y}" if kind == "weekly" else f"{local:%A, %d %B %Y}")
    return ReportData(
        kind=kind,
        date=f"{local:%Y-%m-%d}",
        title=f"{cfg.settings.profile.name} — {'Weekly' if kind == 'weekly' else 'Daily'} report",
        period_label=period,
        generated_at=f"{local:%d %b %Y, %H:%M} IST",
        since=since.isoformat(),
        stories=stories,
        stats=stats,
        top=pick_top(stories, cfg.settings.reports.exec_top_n),
        by_sector=group_by_sector(stories, tax),
        trackers=trackers,
        health=health,
    )


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


def _env():
    from jinja2 import Environment, FileSystemLoader, select_autoescape

    env = Environment(loader=FileSystemLoader(Path(__file__).parent / "templates"),
                      autoescape=select_autoescape(["html"]), trim_blocks=True, lstrip_blocks=True)
    return env


def render(data: ReportData, template: str) -> str:
    return _env().get_template(template).render(r=data)


def write_excel(data: ReportData, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1B2433")

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
        for row in rows:
            ws.append(row)
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    # Summary: state × sector pivot of stories
    summary = wb.active
    summary.title = "Summary"
    summary.append([data.title])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([data.period_label])
    summary.append([])
    sectors = [label for label, _ in data.stats["by_sector"]]
    summary.append(["State / UT", *sectors, "Total"])
    for cell in summary[4]:
        cell.font, cell.fill = header_font, header_fill
    pivot: dict[str, Counter] = defaultdict(Counter)
    for s in data.stories:
        for state in (s.state_names or ["National"]):
            for sector in (s.sector_labels or ["Governance"]):
                pivot[state][sector] += 1
    for state, counts in sorted(pivot.items(), key=lambda kv: (kv[0] != "National", -sum(kv[1].values()))):
        summary.append([state, *[counts.get(sec, 0) for sec in sectors], sum(counts.values())])
    summary.column_dimensions["A"].width = 28
    for i in range(len(sectors) + 1):
        summary.column_dimensions[get_column_letter(i + 2)].width = 18

    columns = ["Score", "Band", "Priority", "Category", "Sector(s)", "State(s)", "Level", "Title", "Source",
               "Outlets", "Language", "First seen (UTC)", "Link"]
    widths = [7, 11, 9, 16, 26, 20, 10, 80, 26, 8, 9, 20, 50]

    def rows_for(stories):
        return [[s.score, s.band, s.priority, s.category_label, ", ".join(s.sector_labels), s.where, s.level,
                 s.title, s.source_name, s.outlets, s.language, s.first_seen, s.url] for s in stories]

    sheet("All stories", columns, rows_for(data.stories), widths)
    for key, label in (("transfer", "Transfers"), ("election", "Elections"), ("court", "Court rulings"),
                       ("policy", "Policy")):
        sheet(label, columns, rows_for(data.trackers[key]), widths)
    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows(min_row=2):
            row[7].alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def telegram_caption(data: ReportData, limit: int = 5) -> str:
    """Short caption for the PDF: headline numbers and the top developments (≤ 1,024 chars)."""
    import html

    s = data.stats
    lines = [f"<b>{'Weekly' if data.kind == 'weekly' else 'Daily'} brief · {html.escape(data.period_label)}</b>",
             f"{s['stories']} stories · {s['core']} core · {s['high_priority']} high priority", ""]
    for i, story in enumerate(data.top[:limit], 1):
        lines.append(f"{i}. {html.escape(story.title[:110])}")
    lines += ["", "Full brief in the PDF. Tap 👍/👎 on the cards below to train."]
    caption = "\n".join(lines)
    return caption if len(caption) <= 1024 else caption[:1000] + "…"


def telegram_text(data: ReportData, limit: int = 8) -> str:
    import html

    s = data.stats
    lines = [f"<b>{html.escape(data.title)}</b>", html.escape(data.period_label), "",
             f"{s['stories']} stories · {s['core']} core · {s['high_priority']} high priority", ""]
    for i, story in enumerate(data.top[:limit], 1):
        extra = f" · {story.outlets} outlets" if story.outlets > 1 else ""
        lines.append(f"{i}. <b>{html.escape(story.title)}</b>\n    {html.escape(story.category_label)} · "
                     f"{html.escape(story.where)}{extra}")
    lines += ["", "Tap the buttons on the cards below to train me. /help"]
    return "\n".join(lines)


def write_report(cfg: AppConfig, data: ReportData, out_root: Path | None = None) -> dict[str, Path]:
    folder = (out_root or cfg.resolve(cfg.settings.reports.output_dir)) / f"{data.kind}-{data.date}"
    folder.mkdir(parents=True, exist_ok=True)
    files = {
        "digest": folder / "analyst-digest.html",
        "brief": folder / "executive-brief.html",
        "excel": folder / "stories.xlsx",
        "telegram": folder / "telegram-brief.txt",
    }
    files["digest"].write_text(render(data, "digest.html"), encoding="utf-8")
    files["brief"].write_text(render(data, "brief.html"), encoding="utf-8")
    write_excel(data, files["excel"])
    files["telegram"].write_text(telegram_text(data), encoding="utf-8")
    from tracker.pdf import html_to_pdf

    pdf = html_to_pdf(files["brief"], folder / f"brief-{data.kind}-{data.date}.pdf")
    if pdf:
        files["pdf"] = pdf
    log.info("report written to %s", folder)
    return files


def send_email(cfg: AppConfig, subject: str, html_body: str, to: list[str], attachments: list[Path] = ()) -> bool:
    sec = cfg.secrets
    if not (sec.smtp_user and sec.smtp_password and to):
        log.info("email not configured; skipping")
        return False
    msg = EmailMessage()
    msg["Subject"], msg["From"], msg["To"] = subject, sec.smtp_user, ", ".join(to)
    msg.set_content("This report is best viewed as HTML.")
    msg.add_alternative(html_body, subtype="html")
    for path in attachments:
        msg.add_attachment(path.read_bytes(), maintype="application", subtype="octet-stream", filename=path.name)
    with smtplib.SMTP(sec.smtp_host, sec.smtp_port, timeout=30) as smtp:
        smtp.starttls()
        smtp.login(sec.smtp_user, sec.smtp_password)
        smtp.send_message(msg)
    return True
